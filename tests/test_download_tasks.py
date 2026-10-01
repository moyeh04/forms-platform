"""Tests for scripts/download_tasks.py, using a small local server in place of Google."""
import http.server
import json
import os
import sys
import tempfile
import threading
import unittest

from openpyxl import load_workbook

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "scripts"))
import download_tasks as dt  # noqa: E402

ID = "a" * 33


class Handler(http.server.BaseHTTPRequestHandler):
    def log_message(self, *args):
        pass

    def do_POST(self):
        body = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
        if body.get("admin", {}).get("pin") != "4321":
            out = {"ok": False, "error": {"code": "bad_pin", "message": "That PIN is not correct."}}
        else:
            out = {"ok": True, "form": {"title": "Tasks", "type": "task_submission"}, "teams": [{"ref": "TK-1", "title": "A", "link": "", "members": []}]}
        data = json.dumps(out).encode()
        self.send_response(200)
        self.send_header("Content-Type", "application/json")
        self.end_headers()
        self.wfile.write(data)

    def do_GET(self):
        if self.path.startswith("/ok"):
            self.send_response(200)
            self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.presentationml.presentation")
            self.end_headers()
            self.wfile.write(b"PPTX-BYTES")
        elif self.path.startswith("/private"):
            self.send_response(200)
            self.send_header("Content-Type", "text/html")
            self.end_headers()
            self.wfile.write(b"<html>Sign in</html>")
        elif self.path.startswith("/scan") and "confirm=" not in self.path:
            self.send_response(200)
            self.send_header("Content-Type", "text/html")
            self.end_headers()
            self.wfile.write(b'<a href="/scan?confirm=T0k3n">Download anyway</a>')
        elif self.path.startswith("/scan"):
            self.send_response(200)
            self.send_header("Content-Type", "application/pdf")
            self.send_header("Content-Disposition", 'attachment; filename="big.pdf"')
            self.end_headers()
            self.wfile.write(b"PDF-BYTES")
        else:
            self.send_response(404)
            self.end_headers()


class ServerCase(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.server = http.server.HTTPServer(("127.0.0.1", 0), Handler)
        cls.base = f"http://127.0.0.1:{cls.server.server_port}"
        threading.Thread(target=cls.server.serve_forever, daemon=True).start()

    @classmethod
    def tearDownClass(cls):
        cls.server.shutdown()


class LinkTests(unittest.TestCase):
    def test_parses_every_kind_of_drive_link(self):
        self.assertEqual(dt.parse_drive_link(f"https://docs.google.com/presentation/d/{ID}/edit"), ("slides", ID))
        self.assertEqual(dt.parse_drive_link(f"https://drive.google.com/file/d/{ID}/view?usp=sharing"), ("file", ID))
        self.assertEqual(dt.parse_drive_link(f"https://drive.google.com/open?id={ID}"), ("file", ID))
        self.assertEqual(dt.parse_drive_link(f"https://drive.google.com/drive/u/0/folders/{ID}"), ("folder", ID))
        self.assertEqual(dt.parse_drive_link("https://example.com/file/d/" + ID), (None, None))
        self.assertEqual(dt.parse_drive_link(""), (None, None))

    def test_slides_are_exported_as_pptx(self):
        url, ext = dt.export_url("slides", ID)
        self.assertTrue(url.endswith("/export/pptx"))
        self.assertEqual(ext, ".pptx")
        self.assertEqual(dt.export_url("folder", ID), (None, None))

    def test_file_names_are_safe(self):
        self.assertEqual(dt.safe_name("Normalization: 1NF/2NF!"), "Normalization_1NF_2NF")
        self.assertEqual(dt.safe_name("نظام"), "task")


class ApiAndDownloadTests(ServerCase):
    def test_fetch_teams_sends_the_pin_and_reads_the_list(self):
        form, teams = dt.fetch_teams(self.base, "4321", "tasks")
        self.assertEqual(form["title"], "Tasks")
        self.assertEqual(len(teams), 1)

    def test_wrong_pin_gives_the_server_message(self):
        with self.assertRaisesRegex(RuntimeError, "not correct"):
            dt.fetch_teams(self.base, "0000", "tasks")

    def test_download_saves_the_bytes_with_the_right_extension(self):
        with tempfile.TemporaryDirectory() as d:
            path, reason = dt.download(self.base + "/ok", os.path.join(d, "001_A"), ".pptx")
            self.assertEqual(reason, "")
            self.assertTrue(path.endswith(".pptx"))
            with open(path, "rb") as f:
                self.assertEqual(f.read(), b"PPTX-BYTES")

    def test_private_links_are_reported_not_saved(self):
        with tempfile.TemporaryDirectory() as d:
            path, reason = dt.download(self.base + "/private", os.path.join(d, "001_A"), ".pptx")
            self.assertIsNone(path)
            self.assertIn("not shared", reason)
            self.assertEqual(os.listdir(d), [])

    def test_virus_scan_confirmation_is_followed_and_the_real_extension_kept(self):
        with tempfile.TemporaryDirectory() as d:
            path, reason = dt.download(self.base + "/scan", os.path.join(d, "001_A"), None)
            self.assertEqual(reason, "")
            self.assertTrue(path.endswith(".pdf"))

    def test_missing_files_report_the_error(self):
        with tempfile.TemporaryDirectory() as d:
            path, reason = dt.download(self.base + "/nope", os.path.join(d, "001_A"), ".pptx")
            self.assertIsNone(path)
            self.assertIn("404", reason)


TEAMS = [
    {"ref": "TK-1", "title": "Alpha", "link": f"https://docs.google.com/presentation/d/{ID}/edit",
     "members": [{"name": "أحمد محمد محمود أحمد", "code": "4230001", "leader": True}, {"name": "سارة خالد حسن علي", "code": "4230002", "leader": False}]},
    {"ref": "TK-2", "title": "Beta", "link": "https://drive.google.com/file/d/" + "b" * 33 + "/view",
     "members": [{"name": "عمر يوسف إبراهيم سعيد", "code": "4230010", "leader": True}, {"name": "منى أشرف كمال فؤاد", "code": "4230011", "leader": False}, {"name": "كريم هشام عادل نصر", "code": "4230012", "leader": False}]},
    {"ref": "TK-3", "title": "Gamma", "link": "https://drive.google.com/file/d/" + "c" * 33 + "/view",
     "members": [{"name": "ياسر محمود عادل سمير", "code": "4230020", "leader": True}]},
]
RESULTS = {
    "TK-1": {"path": "files/001_Alpha.pptx", "reason": ""},
    "TK-2": {"path": "files/002_Beta.pptx", "reason": ""},
    "TK-3": {"path": None, "reason": "The link is not shared with anyone who has the link, or it was deleted."},
}


class WorkbookTests(unittest.TestCase):
    def build(self):
        d = tempfile.mkdtemp()
        path, failed = dt.build_workbook({"title": "Tasks", "type": "task_submission"}, TEAMS, RESULTS, os.path.join(d, "Tasks.xlsx"))
        return load_workbook(path), failed

    def test_layout_matches_the_google_sheet(self):
        wb, failed = self.build()
        ws = wb["Tasks"]
        self.assertEqual([c.value for c in ws[1]], ["Team member name", "Code", "Task"])
        self.assertEqual(ws["A2"].value, "1.  أحمد محمد محمود أحمد  \u2605")
        self.assertEqual(ws["A3"].value, "2.  سارة خالد حسن علي")
        self.assertEqual(ws["B3"].value, "4230002")
        self.assertEqual(ws["A4"].value, "1.  عمر يوسف إبراهيم سعيد  \u2605")
        self.assertEqual(failed, 1)

    def test_each_team_has_one_merged_task_cell_covering_its_members(self):
        ws = self.build()[0]["Tasks"]
        merged = sorted(str(r) for r in ws.merged_cells.ranges)
        self.assertEqual(merged, ["C2:C3", "C4:C6"])
        self.assertEqual(ws["C2"].value, "Alpha")
        self.assertEqual(ws["C4"].value, "Beta")

    def test_task_cells_link_to_the_downloaded_file_with_a_relative_path(self):
        ws = self.build()[0]["Tasks"]
        self.assertEqual(ws["C2"].hyperlink.target, "files/001_Alpha.pptx")
        self.assertEqual(ws["C4"].hyperlink.target, "files/002_Beta.pptx")

    def test_a_failed_download_falls_back_to_the_original_link(self):
        ws = self.build()[0]["Tasks"]
        self.assertTrue(ws["C7"].hyperlink.target.startswith("https://drive.google.com/"))

    def test_neighbouring_teams_alternate_tints_and_rows_inside_a_team_alternate_shades(self):
        ws = self.build()[0]["Tasks"]
        solid = lambda ref: ws[ref].fill.start_color.rgb[-6:]
        self.assertEqual(solid("C2"), "A9CFC6")
        self.assertEqual(solid("C4"), "E6B9B0")
        self.assertEqual(solid("C7"), "A9CFC6")
        self.assertEqual(ws["A2"].fill.start_color.rgb[-6:], "DCEAE4")
        self.assertEqual(ws["A3"].fill.start_color.rgb[-6:], "EEF5F1")
        self.assertEqual(ws["A4"].fill.start_color.rgb[-6:], "F0E4C6")

    def test_heavy_border_wraps_every_team(self):
        ws = self.build()[0]["Tasks"]
        self.assertEqual(ws["A2"].border.top.style, "thick")
        self.assertEqual(ws["A2"].border.left.style, "thick")
        self.assertEqual(ws["A2"].border.bottom.style, "thin")
        self.assertEqual(ws["A3"].border.bottom.style, "thick")
        self.assertEqual(ws["C3"].border.right.style, "thick")
        self.assertEqual(ws["A4"].border.top.style, "thick", "the next team starts with its own heavy line")

    def test_failed_sheet_lists_what_could_not_be_downloaded(self):
        wb = self.build()[0]
        ws = wb["Failed"]
        self.assertEqual([c.value for c in ws[1]], ["Reference", "Task", "Link", "Why it failed"])
        self.assertEqual(ws["A2"].value, "TK-3")
        self.assertIn("not shared", ws["D2"].value)

    def test_no_failed_sheet_when_everything_downloaded(self):
        d = tempfile.mkdtemp()
        ok = {k: {"path": "files/x.pptx", "reason": ""} for k in RESULTS}
        path, failed = dt.build_workbook({"title": "T", "type": "task_submission"}, TEAMS, ok, os.path.join(d, "T.xlsx"))
        self.assertEqual(failed, 0)
        self.assertNotIn("Failed", load_workbook(path).sheetnames)

    def test_projects_use_a_project_heading(self):
        d = tempfile.mkdtemp()
        path, _ = dt.build_workbook({"title": "P", "type": "team_registration"}, TEAMS[:1], {}, os.path.join(d, "P.xlsx"))
        self.assertEqual(load_workbook(path)["Tasks"]["C1"].value, "Project")


class RunTests(ServerCase):
    def test_download_all_names_files_in_order_and_reports_each_failure(self):
        teams = [
            {"ref": "TK-1", "title": "Not a link", "link": "hello", "members": []},
            {"ref": "TK-2", "title": "A folder", "link": f"https://drive.google.com/drive/folders/{ID}", "members": []},
        ]
        with tempfile.TemporaryDirectory() as d:
            out = dt.download_all(teams, d, progress=lambda *_: None)
            self.assertIn("not a Google Drive", out["TK-1"]["reason"])
            self.assertIn("Folders", out["TK-2"]["reason"])
            self.assertTrue(os.path.isdir(os.path.join(d, "files")))

    def test_api_url_is_read_from_the_site_config(self):
        with tempfile.TemporaryDirectory() as d:
            p = os.path.join(d, "config.js")
            with open(p, "w") as f:
                f.write("window.APP_CONFIG = {\n  API_URL: 'https://script.google.com/macros/s/abc/exec'\n};")
            self.assertEqual(dt.api_url_from_config(p), "https://script.google.com/macros/s/abc/exec")
            with open(p, "w") as f:
                f.write("window.APP_CONFIG = { API_URL: 'PASTE_YOUR_WEB_APP_URL_HERE' };")
            self.assertIsNone(dt.api_url_from_config(p))


if __name__ == "__main__":
    unittest.main()
