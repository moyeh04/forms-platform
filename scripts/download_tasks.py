#!/usr/bin/env python3
"""
Download every submitted task as a real file and build an Excel team list.

What you get in the output folder:
    Tasks.xlsx        the team list, same layout as the Google Sheet: one tinted
                      block per team, one merged task cell per team that opens the
                      downloaded file
    files/            each task saved as .pptx (Google Slides are exported)
    Failed sheet      inside Tasks.xlsx, every link that could not be downloaded

Run it from the project folder:
    pip install -r scripts/requirements.txt
    python scripts/download_tasks.py --slug database-tasks-fall-2027

It reads the web app address from assets/js/config.js and asks for the admin PIN.
The links inside Tasks.xlsx are relative, so keep Tasks.xlsx and the files folder
together. They open in desktop Excel; Excel on the web cannot open local files.
"""
import argparse
import getpass
import json
import os
import re
import sys
import urllib.error
import urllib.parse
import urllib.request

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

INK = "3A2A1A"
ON_INK = "F3E6C4"
GRID = "C9B58A"
FAMILIES = [
    {"shades": ("DCEAE4", "EEF5F1"), "solid": "A9CFC6"},
    {"shades": ("F0E4C6", "F8F1DE"), "solid": "E6B9B0"},
]
USER_AGENT = "Mozilla/5.0 (forms-platform task export)"


# ── Google Drive links ─────────────────────────────────────────────

def parse_drive_link(url):
    """Returns (kind, id) for a Drive or Docs link, or (None, None)."""
    url = (url or "").strip()
    if not re.match(r"^https?://(drive|docs)\.google\.com/", url, re.I):
        return None, None
    patterns = [
        (r"/folders/([-\w]{20,})", "folder"),
        (r"/presentation/d/([-\w]{20,})", "slides"),
        (r"/document/d/([-\w]{20,})", "doc"),
        (r"/spreadsheets/d/([-\w]{20,})", "sheet"),
        (r"/file/d/([-\w]{20,})", "file"),
        (r"[?&]id=([-\w]{20,})", "file"),
    ]
    for pattern, kind in patterns:
        m = re.search(pattern, url)
        if m:
            return kind, m.group(1)
    return None, None


def export_url(kind, file_id):
    if kind == "slides":
        return f"https://docs.google.com/presentation/d/{file_id}/export/pptx", ".pptx"
    if kind == "doc":
        return f"https://docs.google.com/document/d/{file_id}/export?format=docx", ".docx"
    if kind == "sheet":
        return f"https://docs.google.com/spreadsheets/d/{file_id}/export?format=xlsx", ".xlsx"
    if kind == "file":
        return f"https://drive.google.com/uc?export=download&id={file_id}", None
    return None, None


def safe_name(text, fallback="task"):
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "_", (text or "").strip()).strip("._-")
    return (cleaned or fallback)[:60]


# ── Talking to the web app ─────────────────────────────────────────

def api_url_from_config(path):
    try:
        with open(path, encoding="utf-8") as f:
            m = re.search(r"API_URL:\s*['\"]([^'\"]+)['\"]", f.read())
    except OSError:
        return None
    if m and not m.group(1).startswith("PASTE_"):
        return m.group(1)
    return None


def call_api(url, payload, timeout=60):
    """POSTs JSON as plain text, the same way the website does."""
    req = urllib.request.Request(
        url,
        data=json.dumps(payload).encode("utf-8"),
        headers={"Content-Type": "text/plain;charset=utf-8", "User-Agent": USER_AGENT},
        method="POST",
    )
    with urllib.request.urlopen(req, timeout=timeout) as resp:
        data = json.loads(resp.read().decode("utf-8"))
    if not data.get("ok"):
        err = data.get("error") or {}
        raise RuntimeError(err.get("message") or "The server refused the request.")
    return data


def fetch_teams(api_url, pin, slug):
    data = call_api(api_url, {"action": "admin.export.tasks", "slug": slug, "admin": {"pin": pin}})
    return data["form"], data["teams"]


# ── Downloading files ──────────────────────────────────────────────

def _extension_from_headers(headers, default):
    disposition = headers.get("Content-Disposition", "")
    m = re.search(r'filename\*?=(?:UTF-8\'\')?"?([^";]+)', disposition)
    if m:
        ext = os.path.splitext(urllib.parse.unquote(m.group(1)))[1].lower()
        if ext:
            return ext
    return default


def download(url, dest_base, default_ext, timeout=120):
    """
    Saves url to dest_base + extension. Returns (path or None, reason).
    Google answers with an HTML page when a file is private or needs a
    virus-scan confirmation; the confirmation is followed once.
    """
    def fetch(u):
        req = urllib.request.Request(u, headers={"User-Agent": USER_AGENT})
        return urllib.request.urlopen(req, timeout=timeout)

    try:
        resp = fetch(url)
        body = resp.read()
        headers = resp.headers
        if "text/html" in headers.get("Content-Type", ""):
            m = re.search(rb'confirm=([0-9A-Za-z_-]+)', body)
            if m:
                sep = "&" if "?" in url else "?"
                resp = fetch(url + sep + "confirm=" + m.group(1).decode())
                body = resp.read()
                headers = resp.headers
            if "text/html" in headers.get("Content-Type", ""):
                return None, "The link is not shared with anyone who has the link, or it was deleted."
        ext = _extension_from_headers(headers, default_ext or ".pptx")
        path = dest_base + ext
        with open(path, "wb") as f:
            f.write(body)
        return path, ""
    except urllib.error.HTTPError as e:
        return None, f"Google answered with error {e.code}."
    except (urllib.error.URLError, TimeoutError, OSError) as e:
        return None, f"Could not download: {e}"


def download_all(teams, out_dir, progress=print):
    """Returns {ref: {"path": relative path or None, "reason": text}}."""
    files_dir = os.path.join(out_dir, "files")
    os.makedirs(files_dir, exist_ok=True)
    results = {}
    for i, team in enumerate(teams, start=1):
        kind, file_id = parse_drive_link(team.get("link"))
        if kind is None:
            results[team["ref"]] = {"path": None, "reason": "The link is not a Google Drive or Slides link."}
            progress(f"[{i}/{len(teams)}] {team['title']}: skipped (not a Drive link)")
            continue
        url, ext = export_url(kind, file_id)
        if url is None:
            results[team["ref"]] = {"path": None, "reason": "Folders cannot be downloaded as one file."}
            progress(f"[{i}/{len(teams)}] {team['title']}: skipped (folder)")
            continue
        base = os.path.join(files_dir, f"{i:03d}_{safe_name(team['title'])}")
        path, reason = download(url, base, ext)
        if path:
            results[team["ref"]] = {"path": os.path.relpath(path, out_dir).replace(os.sep, "/"), "reason": ""}
            progress(f"[{i}/{len(teams)}] {team['title']}: saved")
        else:
            results[team["ref"]] = {"path": None, "reason": reason}
            progress(f"[{i}/{len(teams)}] {team['title']}: FAILED - {reason}")
    return results


# ── The Excel team list ────────────────────────────────────────────

def _fill(color):
    return PatternFill("solid", start_color=color, end_color=color)


def build_workbook(form, teams, results, out_path):
    """One tinted block per team, a heavy border around it, one merged task cell."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Tasks"
    task_label = "Task" if form.get("type") == "task_submission" else "Project"

    thin = Side(style="thin", color=GRID)
    thick = Side(style="thick", color=INK)
    base_font = "Arial"

    for col, text in enumerate(["Team member name", "Code", task_label], start=1):
        c = ws.cell(row=1, column=col, value=text)
        c.font = Font(name=base_font, bold=True, size=12, color=ON_INK)
        c.fill = _fill(INK)
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 42

    row = 2
    for g, team in enumerate(teams):
        family = FAMILIES[g % 2]
        members = team["members"] or [{"name": "", "code": "", "leader": False}]
        first, last = row, row + len(members) - 1

        for i, m in enumerate(members):
            shade = _fill(family["shades"][i % 2])
            name = f"{i + 1}.  {m['name']}" + ("  \u2605" if m.get("leader") else "")
            a = ws.cell(row=row, column=1, value=name)
            a.font = Font(name=base_font, bold=True, size=11)
            a.alignment = Alignment(horizontal="left", vertical="center", readingOrder=2)
            b = ws.cell(row=row, column=2, value=m["code"])
            b.font = Font(name=base_font, bold=True, size=11)
            b.alignment = Alignment(horizontal="center", vertical="center")
            for cell in (a, b):
                cell.fill = shade
            ws.row_dimensions[row].height = 24
            row += 1

        title = team.get("title") or "-"
        task = ws.cell(row=first, column=3, value=title)
        if last > first:
            ws.merge_cells(start_row=first, start_column=3, end_row=last, end_column=3)
        result = results.get(team["ref"], {})
        target = result.get("path") or team.get("link")
        if target:
            task.hyperlink = target
        task.font = Font(name=base_font, bold=True, size=12, color=INK, underline="single" if target else None)
        task.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for r in range(first, last + 1):
            ws.cell(row=r, column=3).fill = _fill(family["solid"])

        # Thin grid inside, heavy frame around the whole team.
        for r in range(first, last + 1):
            for c in range(1, 4):
                ws.cell(row=r, column=c).border = Border(
                    left=thick if c == 1 else thin,
                    right=thick if c == 3 else thin,
                    top=thick if r == first else thin,
                    bottom=thick if r == last else thin,
                )

    failed = [(t, results[t["ref"]]["reason"]) for t in teams if t["ref"] in results and not results[t["ref"]]["path"]]
    if failed:
        fs = wb.create_sheet("Failed")
        for col, text in enumerate(["Reference", task_label, "Link", "Why it failed"], start=1):
            c = fs.cell(row=1, column=col, value=text)
            c.font = Font(name=base_font, bold=True, color=ON_INK)
            c.fill = _fill(INK)
        for i, (t, reason) in enumerate(failed, start=2):
            fs.cell(row=i, column=1, value=t["ref"])
            fs.cell(row=i, column=2, value=t.get("title") or "")
            link_cell = fs.cell(row=i, column=3, value=t.get("link") or "")
            if t.get("link"):
                link_cell.hyperlink = t["link"]
            fs.cell(row=i, column=4, value=reason)
        for col, width in zip("ABCD", (18, 30, 60, 60)):
            fs.column_dimensions[col].width = width

    wb.save(out_path)
    return out_path, len(failed)


# ── Command line ───────────────────────────────────────────────────

def main(argv=None):
    here = os.path.dirname(os.path.abspath(__file__))
    root = os.path.dirname(here)
    ap = argparse.ArgumentParser(description="Download task files and build an Excel team list.")
    ap.add_argument("--slug", required=True, help="the form's link name (the part after ?f=)")
    ap.add_argument("--api", help="web app address (default: read from assets/js/config.js)")
    ap.add_argument("--pin", help="admin PIN (default: FORMS_ADMIN_PIN or asked)")
    ap.add_argument("--out", help="output folder (default: Tasks_Export/<slug>)")
    ap.add_argument("--no-download", action="store_true", help="only build the list, keep the original links")
    args = ap.parse_args(argv)

    api_url = args.api or os.environ.get("FORMS_API_URL") or api_url_from_config(os.path.join(root, "assets", "js", "config.js"))
    if not api_url:
        sys.exit("Put the web app address in assets/js/config.js or pass --api.")
    pin = args.pin or os.environ.get("FORMS_ADMIN_PIN") or getpass.getpass("Admin PIN: ")

    try:
        form, teams = fetch_teams(api_url, pin, args.slug)
    except (RuntimeError, urllib.error.URLError) as e:
        sys.exit(f"Could not get the team list: {e}")
    if not teams:
        sys.exit("There are no submissions to export yet.")

    out_dir = args.out or os.path.join(root, "Tasks_Export", args.slug)
    os.makedirs(out_dir, exist_ok=True)
    print(f"{form['title']}: {len(teams)} teams")
    results = {} if args.no_download else download_all(teams, out_dir)
    path, failed = build_workbook(form, teams, results, os.path.join(out_dir, "Tasks.xlsx"))
    print(f"\nSaved {path}")
    if failed:
        print(f"{failed} link(s) could not be downloaded. They are listed on the Failed sheet.")


if __name__ == "__main__":
    main()
